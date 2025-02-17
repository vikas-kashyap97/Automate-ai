import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import docx

class Teleprompter:
    def __init__(self, master):
        self.master = master
        master.title("Modern Teleprompter")
        master.geometry("800x600")
        master.resizable(False, False)  # Disable window resizing
        
        style = ttk.Style()
        style.theme_use('clam')
        
        self.main_frame = ttk.Frame(master)
        self.main_frame.pack(fill=tk.BOTH, expand=True)
        
        # Control frame at the top
        self.control_frame = ttk.Frame(self.main_frame)
        self.control_frame.pack(side=tk.TOP, fill=tk.X, padx=10, pady=10)
        
        self.load_button = ttk.Button(self.control_frame, text="Load Script", command=self.load_script)
        self.load_button.grid(row=0, column=0, padx=5, pady=5, sticky="ew")
        
        self.start_button = ttk.Button(self.control_frame, text="Start", command=self.start_scrolling)
        self.start_button.grid(row=0, column=1, padx=5, pady=5, sticky="ew")
        
        self.stop_button = ttk.Button(self.control_frame, text="Stop", command=self.stop_scrolling)
        self.stop_button.grid(row=0, column=2, padx=5, pady=5, sticky="ew")
        
        self.restart_button = ttk.Button(self.control_frame, text="Restart", command=self.restart_scrolling)
        self.restart_button.grid(row=0, column=3, padx=5, pady=5, sticky="ew")
        
        self.speed_label = ttk.Label(self.control_frame, text="Speed:")
        self.speed_label.grid(row=1, column=0, padx=5, pady=5, sticky="w")
        
        self.speed_scale = ttk.Scale(self.control_frame, from_=0.1, to=50, orient=tk.HORIZONTAL)
        self.speed_scale.set(5)
        self.speed_scale.grid(row=1, column=1, columnspan=3, padx=5, pady=5, sticky="ew")
        
        self.font_size_label = ttk.Label(self.control_frame, text="Font Size:")
        self.font_size_label.grid(row=2, column=0, padx=5, pady=5, sticky="w")
        
        self.font_size_scale = ttk.Scale(self.control_frame, from_=12, to=72, orient=tk.HORIZONTAL, command=self.change_font_size)
        self.font_size_scale.set(24)
        self.font_size_scale.grid(row=2, column=1, columnspan=3, padx=5, pady=5, sticky="ew")
        
        self.control_frame.grid_columnconfigure((0, 1, 2, 3), weight=1)
        
        # Text frame below the control frame
        self.text_frame = ttk.Frame(self.main_frame)
        self.text_frame.pack(fill=tk.BOTH, expand=True)
        
        self.text = tk.Text(self.text_frame, wrap=tk.WORD, font=("Arial", 24), bg="black", fg="white")
        self.text.pack(fill=tk.BOTH, expand=True)
        
        self.scrollbar = ttk.Scrollbar(self.text_frame, orient="vertical", command=self.text.yview)
        self.scrollbar.pack(side="right", fill="y")
        self.text.configure(yscrollcommand=self.scrollbar.set)
        
        self.scrolling = False
        self.scroll_position = 0.0
        
    def load_script(self):
        file_path = filedialog.askopenfilename(filetypes=[("Word Document", "*.docx"), ("Text File", "*.txt")])
        if file_path:
            try:
                if file_path.endswith('.docx'):
                    doc = docx.Document(file_path)
                    full_text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
                else:  # Assume it's a .txt file
                    with open(file_path, 'r', encoding='utf-8') as file:
                        full_text = file.read()
                self.text.delete(1.0, tk.END)
                self.text.insert(tk.END, full_text)
                self.scroll_position = 0.0
                self.text.yview_moveto(0)
            except Exception as e:
                messagebox.showerror("Error", f"Failed to load the document: {str(e)}")
        
    def start_scrolling(self):
        self.scrolling = True
        self.scroll_text()
        
    def stop_scrolling(self):
        self.scrolling = False
        
    def restart_scrolling(self):
        self.scroll_position = 0.0
        self.text.yview_moveto(0)
        self.start_scrolling()
        
    def scroll_text(self):
        if self.scrolling:
            self.scroll_position += 0.0001 * self.speed_scale.get()
            self.text.yview_moveto(self.scroll_position)
            if self.scroll_position >= 1.0:
                self.stop_scrolling()
                return
            self.master.after(20, self.scroll_text)
        
    def change_font_size(self, size):
        new_size = int(float(size))
        self.text.configure(font=("Arial", new_size))
        # Adjust text widget height to maintain visibility of control frame
        approx_lines = 500 // new_size  # Reduced from 600 to account for control frame
        self.text.configure(height=approx_lines)

root = tk.Tk()
teleprompter = Teleprompter(root)
root.mainloop()